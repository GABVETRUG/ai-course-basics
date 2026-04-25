#!/usr/bin/env python3
"""
Certificato v2 — "Dark & Gold": sfondo scuro, accenti oro, stile premium.
Genera versione IT e EN.
"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# ── Palette "Dark & Gold" ───────────────────────────────
DARK      = RGBColor(0x0F, 0x11, 0x16)   # sfondo principale
DARK_2    = RGBColor(0x18, 0x1B, 0x22)   # card sfondo
DARK_3    = RGBColor(0x22, 0x26, 0x2F)   # bordi sottili
GOLD      = RGBColor(0xD4, 0xAF, 0x37)   # oro classico
GOLD_L    = RGBColor(0xE8, 0xCF, 0x7A)   # oro chiaro
GOLD_D    = RGBColor(0xA0, 0x82, 0x2A)   # oro scuro
CREAM     = RGBColor(0xF5, 0xF0, 0xE8)   # testo principale
CREAM_D   = RGBColor(0xB8, 0xB0, 0xA0)   # testo secondario
CREAM_DD  = RGBColor(0x78, 0x72, 0x68)   # testo dim
BIANCO    = RGBColor(0xFF, 0xFF, 0xFF)

F_DISP = "Georgia"
F_UI   = "Calibri"
F_MONO = "Consolas"

SW = Inches(13.333)
SH = Inches(7.5)


def _kill_line(shape):
    spPr = shape._element.spPr
    ln = spPr.find(qn('a:ln'))
    if ln is None:
        ln = etree.SubElement(spPr, qn('a:ln'))
    else:
        for child in list(ln):
            ln.remove(child)
    ln.set('w', '0')
    etree.SubElement(ln, qn('a:noFill'))


def _rect(slide, l, t, w, h, color, line_color=None, line_w=None, radius=None):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(st, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color
    if line_color:
        s.line.color.rgb = line_color; s.line.width = Pt(line_w or 1)
    else:
        _kill_line(s)
    return s


def _tb(slide, l, t, w, h, txt, sz=18, bold=False, col=CREAM,
        align=PP_ALIGN.LEFT, font=F_UI, italic=False):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    _kill_line(tb)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = txt; p.font.size = Pt(sz); p.font.bold = bold
    p.font.color.rgb = col; p.font.name = font
    p.font.italic = italic; p.alignment = align
    for run in p.runs:
        rPr = run._r.find(qn('a:rPr'))
        if rPr is not None:
            rPr.set('noProof', '1')
    return tf


def _diamond(slide, cx, cy, size, color):
    """Rombo decorativo (quadrato ruotato 45deg)."""
    s = slide.shapes.add_shape(MSO_SHAPE.DIAMOND,
        Inches(cx - size/2), Inches(cy - size/2), Inches(size), Inches(size))
    s.fill.solid(); s.fill.fore_color.rgb = color; _kill_line(s)


def _corner_L(slide, x, y, length, color, flip_h=False, flip_v=False):
    """Angolo decorativo a L (due linee sottili perpendicolari)."""
    lw, lh = length, 0.02
    if flip_h:
        _rect(slide, x - length, y, length, lh, color)
    else:
        _rect(slide, x, y, length, lh, color)
    if flip_v:
        _rect(slide, x if not flip_h else x - lw + lh, y - length, lh, length, color)
    else:
        _rect(slide, x if not flip_h else x - lh, y, lh, length, color)


def genera_certificato_v2(lang="it"):
    prs = Presentation()
    prs.slide_width = SW; prs.slide_height = SH
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # ── Sfondo scuro ────────────────────────────────────
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = DARK

    # ── Cornice oro doppia ──────────────────────────────
    _rect(s, 0.4, 0.3, 12.53, 6.9, DARK, line_color=GOLD_D, line_w=0.5)
    _rect(s, 0.55, 0.45, 12.23, 6.6, DARK, line_color=GOLD, line_w=1.0)

    # ── Angoli decorativi oro (4 angoli a L) ────────────
    corner_len = 0.6
    # Top-left
    _corner_L(s, 0.7, 0.6, corner_len, GOLD)
    # Top-right
    _corner_L(s, 12.63, 0.6, corner_len, GOLD, flip_h=True)
    # Bottom-left
    _corner_L(s, 0.7, 6.9, corner_len, GOLD, flip_v=True)
    # Bottom-right
    _corner_L(s, 12.63, 6.9, corner_len, GOLD, flip_h=True, flip_v=True)

    # ── Rombi decorativi centrali (sopra e sotto il nome) ─
    center_x = 6.666
    _diamond(s, center_x, 1.15, 0.12, GOLD)
    _diamond(s, center_x - 0.25, 1.15, 0.06, GOLD_D)
    _diamond(s, center_x + 0.25, 1.15, 0.06, GOLD_D)

    # ── Contenuti ───────────────────────────────────────
    if lang == "it":
        eyebrow = "CERTIFICATO DI COMPLETAMENTO"
        pre_title = "Si attesta che"
        participant = "[NOME E COGNOME]"
        post_title = "ha completato con successo il corso"
        course = "AI Basics"
        course_sub = "Strumenti di Intelligenza Artificiale per Tutti"
        duration = "20 ore  ·  8 moduli  ·  Teoria + Pratica"
        skills_label = "COMPETENZE CERTIFICATE"
        skills = [
            "Uso professionale di ChatGPT, Claude e modelli open source",
            "Prompt engineering con tecniche strutturate",
            "Integrazione dell'AI nei workflow lavorativi",
            "Uso responsabile e consapevole dell'AI",
        ]
        score_label = "VALUTAZIONE"
        score_fields = "Progetto: ___ / 25     Quiz: ___ / 20"
        date_label = "Data"
        date_val = "____________________"
        sign_label = "Firma del Formatore"
        sign_val = "____________________"
    else:
        eyebrow = "CERTIFICATE OF COMPLETION"
        pre_title = "This certifies that"
        participant = "[FULL NAME]"
        post_title = "has successfully completed the course"
        course = "AI Basics"
        course_sub = "Artificial Intelligence Tools for Everyone"
        duration = "20 hours  ·  8 modules  ·  Theory + Practice"
        skills_label = "CERTIFIED SKILLS"
        skills = [
            "Professional use of ChatGPT, Claude and open source models",
            "Prompt engineering with structured techniques",
            "AI integration into daily work workflows",
            "Responsible and informed use of AI",
        ]
        score_label = "ASSESSMENT"
        score_fields = "Project: ___ / 25     Quiz: ___ / 20"
        date_label = "Date"
        date_val = "____________________"
        sign_label = "Instructor Signature"
        sign_val = "____________________"

    # ── Eyebrow ─────────────────────────────────────────
    _tb(s, 1.0, 0.65, 11.3, 0.3, eyebrow,
        sz=12, col=GOLD, font=F_MONO, bold=True, align=PP_ALIGN.CENTER)

    # ── Pre-titolo ──────────────────────────────────────
    _tb(s, 1.0, 1.4, 11.3, 0.35, pre_title,
        sz=18, col=CREAM_D, font=F_UI, italic=True, align=PP_ALIGN.CENTER)

    # ── Nome partecipante ───────────────────────────────
    _tb(s, 1.0, 1.8, 11.3, 0.9, participant,
        sz=56, bold=True, col=GOLD_L, font=F_DISP, italic=True,
        align=PP_ALIGN.CENTER)

    # ── Linea oro sotto il nome ─────────────────────────
    _rect(s, 3.5, 2.7, 6.3, 0.025, GOLD)
    # Rombo centrale sulla linea
    _diamond(s, center_x, 2.71, 0.1, GOLD)

    # ── Post-titolo ─────────────────────────────────────
    _tb(s, 1.0, 2.9, 11.3, 0.35, post_title,
        sz=18, col=CREAM_D, font=F_UI, italic=True, align=PP_ALIGN.CENTER)

    # ── Nome corso ──────────────────────────────────────
    _tb(s, 1.0, 3.3, 11.3, 0.7, course,
        sz=44, bold=True, col=CREAM, font=F_DISP, align=PP_ALIGN.CENTER)

    # ── Sottotitolo corso ───────────────────────────────
    _tb(s, 1.0, 3.9, 11.3, 0.5, course_sub,
        sz=22, col=CREAM_D, font=F_DISP, italic=True, align=PP_ALIGN.CENTER)

    # ── Durata ──────────────────────────────────────────
    _tb(s, 1.0, 4.35, 11.3, 0.3, duration,
        sz=12, col=CREAM_DD, font=F_MONO, align=PP_ALIGN.CENTER)

    # ── Sezione competenze ──────────────────────────────
    # Linee oro decorative sopra e sotto
    _rect(s, 3.0, 4.8, 7.3, 0.015, GOLD_D)

    _tb(s, 1.0, 4.9, 11.3, 0.25, skills_label,
        sz=10, col=GOLD, font=F_MONO, bold=True, align=PP_ALIGN.CENTER)

    # 4 competenze in lista verticale centrata
    for i, skill in enumerate(skills):
        y = 5.2 + i * 0.28
        # Rombo marker
        _diamond(s, 3.8, y + 0.1, 0.06, GOLD)
        # Testo
        _tb(s, 4.0, y, 6.5, 0.25, skill,
            sz=12, col=CREAM_D, font=F_UI)

    _rect(s, 3.0, 6.4, 7.3, 0.015, GOLD_D)

    # ── Valutazione + Data + Firma ──────────────────────
    # Valutazione (sinistra)
    _tb(s, 1.2, 6.55, 3.5, 0.18, score_label,
        sz=9, col=GOLD, font=F_MONO, bold=True)
    _tb(s, 1.2, 6.72, 3.5, 0.2, score_fields,
        sz=12, col=CREAM_D, font=F_MONO)

    # Data (centro)
    _tb(s, 5.2, 6.55, 3, 0.18, date_label,
        sz=9, col=GOLD, font=F_MONO, bold=True, align=PP_ALIGN.CENTER)
    _tb(s, 5.2, 6.72, 3, 0.2, date_val,
        sz=13, col=CREAM_D, font=F_UI, align=PP_ALIGN.CENTER)

    # Firma (destra)
    _tb(s, 8.8, 6.55, 3.9, 0.18, sign_label,
        sz=9, col=GOLD, font=F_MONO, bold=True, align=PP_ALIGN.CENTER)
    _tb(s, 8.8, 6.72, 3.9, 0.2, sign_val,
        sz=13, col=CREAM_D, font=F_UI, align=PP_ALIGN.CENTER)

    return prs


# ── Genera entrambe le versioni ─────────────────────────
if __name__ == "__main__":
    base = os.path.dirname(os.path.abspath(__file__))

    # IT
    prs_it = genera_certificato_v2("it")
    out_it = os.path.join(base, "materiale_trasversale",
                          "Certificato_AI_Basics_DarkGold.pptx")
    prs_it.save(out_it)
    print(f"  OK  {out_it}")

    # EN
    prs_en = genera_certificato_v2("en")
    out_en = os.path.join(base, "..", "ai_course_basics_en", "cross_cutting",
                          "Certificate_AI_Basics_DarkGold.pptx")
    os.makedirs(os.path.dirname(out_en), exist_ok=True)
    prs_en.save(out_en)
    print(f"  OK  {out_en}")

    print("  Certificati Dark & Gold generati!")
