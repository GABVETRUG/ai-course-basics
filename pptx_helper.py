"""
pptx_helper.py — Stile "paper + vibrant accents" ispirato a Oracle PL/SQL Basics v2.
v3 — fix: accent line overlap titoli lunghi, spell-check underlines, bordi shape.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# ── Palette ─────────────────────────────────────────────
PAPER     = RGBColor(0xFA, 0xFA, 0xF7)
PAPER_2   = RGBColor(0xF2, 0xF1, 0xEC)
INK       = RGBColor(0x0B, 0x0B, 0x0C)
INK_2     = RGBColor(0x2A, 0x2A, 0x2D)
INK_DIM   = RGBColor(0x6B, 0x6B, 0x72)
INK_FAINT = RGBColor(0xB5, 0xB5, 0xBC)
RULE      = RGBColor(0xE6, 0xE5, 0xDF)
BIANCO    = RGBColor(0xFF, 0xFF, 0xFF)

LIME      = RGBColor(0xC6, 0xFF, 0x3D)
TANGERINE = RGBColor(0xFF, 0x5A, 0x1F)
COBALT    = RGBColor(0x1E, 0x3C, 0xFF)
MAGENTA   = RGBColor(0xFF, 0x2E, 0x88)
AMBER     = RGBColor(0xFF, 0xB8, 0x00)
MINT      = RGBColor(0x00, 0xC2, 0x8A)

ACCENTS = [TANGERINE, COBALT, MAGENTA, MINT, AMBER, LIME]

# ── Font ────────────────────────────────────────────────
F_DISP = "Georgia"
F_UI   = "Calibri"
F_MONO = "Consolas"

SW = Inches(13.333)
SH = Inches(7.5)

# Layout
ML = 1.1          # margine sinistro
CW = 11.133       # larghezza contenuto
Y_STRIP_TOP = 0.5
Y_STRIP_BOT = 6.9


# ── Utilita interne ─────────────────────────────────────
def crea_presentazione():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    return prs


def _bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _kill_line(shape):
    """Rimuove completamente il bordo di una shape via XML (<a:ln w='0'><a:noFill/>)."""
    spPr = shape._element.spPr
    ln = spPr.find(qn('a:ln'))
    if ln is None:
        ln = etree.SubElement(spPr, qn('a:ln'))
    else:
        for child in list(ln):
            ln.remove(child)
    ln.set('w', '0')
    etree.SubElement(ln, qn('a:noFill'))


def _rect(slide, l, t, w, h, color, line_color=None, line_w=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if line_color:
        s.line.color.rgb = line_color
        s.line.width = Pt(line_w or 1)
    else:
        _kill_line(s)
    return s


def _dot(slide, l, t, color, size=0.1):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                               Inches(l), Inches(t), Inches(size), Inches(size))
    c.fill.solid()
    c.fill.fore_color.rgb = color
    _kill_line(c)
    return c


def _no_proof(paragraph):
    """Disabilita spell-check su un paragrafo (evita sottolineature rosse)."""
    for run in paragraph.runs:
        rPr = run._r.find(qn('a:rPr'))
        if rPr is not None:
            rPr.set('lang', 'it-IT')
            rPr.set('noProof', '1')
            rPr.set('dirty', '0')


def _tb(slide, l, t, w, h, txt, sz=18, bold=False, col=INK,
        align=PP_ALIGN.LEFT, font=F_UI, italic=False):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    _kill_line(tb)  # rimuove bordo default (visibile su sfondo scuro)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = txt
    p.font.size = Pt(sz)
    p.font.bold = bold
    p.font.color.rgb = col
    p.font.name = font
    p.font.italic = italic
    p.alignment = align
    _no_proof(p)
    return tf


def _bullets(tf, items, sz=20, col=INK_2, font=F_UI, spacing=8):
    for i, txt in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if txt.strip() == "":
            p.text = ""
            p.font.size = Pt(8)
            p.space_after = Pt(2)
            continue
        stripped = txt.lstrip()
        indent = len(txt) - len(stripped)
        if indent > 0:
            p.text = f"     {stripped}"
            p.font.color.rgb = INK_DIM
        else:
            p.text = f" ·  {txt}"
            p.font.color.rgb = col
        p.font.size = Pt(sz)
        p.font.name = font
        p.space_after = Pt(spacing)
        _no_proof(p)


def _strip_top(slide, left, right="", dot_color=TANGERINE, text_col=INK_DIM):
    _dot(slide, ML, Y_STRIP_TOP + 0.06, dot_color)
    _tb(slide, ML + 0.2, Y_STRIP_TOP, 6, 0.3, left,
        sz=12, col=text_col, font=F_MONO)
    if right:
        _tb(slide, ML + 6, Y_STRIP_TOP, CW - 6, 0.3, right,
            sz=12, col=text_col, font=F_MONO, align=PP_ALIGN.RIGHT)


def _accent_for(num):
    return ACCENTS[(num - 1) % len(ACCENTS)]


# ── Slide pubbliche ────────────────────────────────────

def title_slide(prs, titolo, sottotitolo="", footer="Corso AI Basics"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, PAPER)

    _strip_top(s, "AI Basics  ·  Corso completo", "2026", TANGERINE)

    # Pill decorativi
    pills = [
        (ML, 1.1, "ChatGPT", LIME, INK),
        (ML + 1.55, 1.1, "Claude", TANGERINE, BIANCO),
        (ML + 2.85, 1.1, "Open Source", COBALT, BIANCO),
        (ML + 4.45, 1.1, "Prompt Eng.", MAGENTA, BIANCO),
    ]
    for (pl, pt, ptxt, pbg, ptc) in pills:
        _rect(s, pl, pt, 1.4, 0.3, pbg, radius=True)
        _tb(s, pl + 0.05, pt + 0.02, 1.3, 0.26, ptxt,
            sz=11, col=ptc, font=F_MONO, bold=True, align=PP_ALIGN.CENTER)

    # Titolo GRANDE
    _tb(s, ML, 1.8, CW, 1.8, titolo,
        sz=120, bold=True, col=INK, font=F_DISP)

    # Linea accento
    _rect(s, ML, 3.7, 2.5, 0.06, TANGERINE)

    # Sottotitolo serif italic
    if sottotitolo:
        _tb(s, ML, 4.0, CW, 1.0, sottotitolo,
            sz=44, col=INK_2, font=F_DISP, italic=True)

    # Footer
    _tb(s, ML, Y_STRIP_BOT, CW, 0.3, footer,
        sz=12, col=INK_DIM, font=F_MONO)
    return s


def section_slide(prs, num, titolo, sottotitolo=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, INK)
    accent = _accent_for(num)

    # Numero grande
    _tb(s, ML, 1.2, 4, 1.5, f"{num:02d}",
        sz=120, col=accent, font=F_DISP, italic=True)

    # Titolo — 2.2" di altezza per titoli fino a 3 righe a 48pt
    _tb(s, ML, 2.7, CW, 2.2, titolo,
        sz=48, bold=True, col=BIANCO, font=F_DISP)

    # Linea accento — posizionata a y=5.3, sotto anche titoli di 3 righe
    _rect(s, ML, 5.3, 1.5, 0.05, accent)

    # Sottotitolo
    if sottotitolo:
        _tb(s, ML, 5.55, CW, 0.7, sottotitolo,
            sz=20, col=INK_FAINT, font=F_UI)

    # Label bottom
    _tb(s, ML, Y_STRIP_BOT, 6, 0.3, f"Sezione {num:02d}",
        sz=12, col=INK_DIM, font=F_MONO)
    return s


def content_slide(prs, titolo, bullets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, PAPER)

    # Eyebrow: label corso (NON il titolo — evita duplicazione)
    _strip_top(s, "Corso AI Basics", dot_color=COBALT)

    # Titolo display
    _tb(s, ML, 1.1, CW, 0.7, titolo,
        sz=34, bold=True, col=INK, font=F_DISP)

    # Riga sottile
    _rect(s, ML, 1.75, CW, 0.012, RULE)

    # Bullets — occupano tutto lo spazio disponibile
    tb = s.shapes.add_textbox(Inches(ML), Inches(1.95),
                              Inches(CW), Inches(4.7))
    _kill_line(tb)
    tf = tb.text_frame
    tf.word_wrap = True
    _bullets(tf, bullets, sz=19, col=INK_2, spacing=7)

    return s


def two_col_slide(prs, titolo, lt, lb, rt, rb):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, PAPER)

    _strip_top(s, "Corso AI Basics", dot_color=MAGENTA)

    # Titolo
    _tb(s, ML, 1.1, CW, 0.6, titolo,
        sz=34, bold=True, col=INK, font=F_DISP)
    _rect(s, ML, 1.7, CW, 0.012, RULE)

    col_w = 5.3
    col_gap = 0.5
    y_top = 1.9
    card_h = 4.5

    # Card sinistra
    _rect(s, ML, y_top, col_w, card_h, BIANCO, line_color=RULE, radius=True)
    _tb(s, ML + 0.3, y_top + 0.2, col_w - 0.6, 0.35, lt,
        sz=22, bold=True, col=TANGERINE, font=F_DISP)
    # Linea colorata sotto titolo card
    _rect(s, ML + 0.3, y_top + 0.6, 1.0, 0.035, TANGERINE)
    tb_l = s.shapes.add_textbox(Inches(ML + 0.3), Inches(y_top + 0.8),
                                Inches(col_w - 0.6), Inches(card_h - 1.1))
    _kill_line(tb_l)
    tf_l = tb_l.text_frame; tf_l.word_wrap = True
    _bullets(tf_l, lb, sz=16, col=INK_2, spacing=6)

    # Card destra
    x_r = ML + col_w + col_gap
    _rect(s, x_r, y_top, col_w, card_h, BIANCO, line_color=RULE, radius=True)
    _tb(s, x_r + 0.3, y_top + 0.2, col_w - 0.6, 0.35, rt,
        sz=22, bold=True, col=COBALT, font=F_DISP)
    _rect(s, x_r + 0.3, y_top + 0.6, 1.0, 0.035, COBALT)
    tb_r = s.shapes.add_textbox(Inches(x_r + 0.3), Inches(y_top + 0.8),
                                Inches(col_w - 0.6), Inches(card_h - 1.1))
    _kill_line(tb_r)
    tf_r = tb_r.text_frame; tf_r.word_wrap = True
    _bullets(tf_r, rb, sz=16, col=INK_2, spacing=6)

    return s


def table_slide(prs, titolo, headers, rows):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, PAPER)

    _strip_top(s, "Corso AI Basics", dot_color=AMBER)
    _tb(s, ML, 1.1, CW, 0.6, titolo,
        sz=34, bold=True, col=INK, font=F_DISP)
    _rect(s, ML, 1.7, CW, 0.012, RULE)

    nr, nc = len(rows) + 1, len(headers)
    tw = Inches(CW)
    row_h = min(0.48, 4.8 / nr)
    th = Inches(row_h) * nr

    tbl = s.shapes.add_table(nr, nc, Inches(ML), Inches(1.9),
                             tw, th).table
    cw = tw // nc
    for i in range(nc):
        tbl.columns[i].width = cw

    # Header
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i); cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(14); p.font.bold = True
            p.font.color.rgb = BIANCO; p.font.name = F_UI
        cell.fill.solid(); cell.fill.fore_color.rgb = INK

    # Rows
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = tbl.cell(r + 1, c); cell.text = str(v)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(13); p.font.color.rgb = INK_2
                p.font.name = F_UI
            cell.fill.solid()
            cell.fill.fore_color.rgb = BIANCO if r % 2 == 0 else PAPER_2
    return s


def highlight_slide(prs, label, testo, desc=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, PAPER)

    # Eyebrow con label (senza duplicazione — usa label generico)
    _strip_top(s, "Concetto chiave", dot_color=LIME)

    # Label mono
    _tb(s, ML, 1.2, CW, 0.3, label,
        sz=13, col=INK_DIM, font=F_MONO, bold=True)

    # Testo GRANDE serif — impatto visivo
    _tb(s, ML, 1.8, CW, 3.0, testo,
        sz=44, col=INK, font=F_DISP, italic=False)

    # Linea accento
    _rect(s, ML, 5.1, 1.5, 0.05, LIME)

    # Descrizione
    if desc:
        _tb(s, ML, 5.35, CW, 1.0, desc,
            sz=18, col=INK_DIM, font=F_UI)

    return s


def recap_slide(prs, punti):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, INK)

    # Label accento
    _tb(s, ML, 0.7, 3, 0.3, "RIEPILOGO",
        sz=13, col=LIME, font=F_MONO, bold=True)

    # Titolo
    _tb(s, ML, 1.05, CW, 0.7, "Concetti chiave",
        sz=42, bold=True, col=BIANCO, font=F_DISP, italic=True)

    # Linea accento
    _rect(s, ML, 1.85, 1.5, 0.05, LIME)

    # Punti — riempiono bene lo spazio
    tb = s.shapes.add_textbox(Inches(ML), Inches(2.2), Inches(CW), Inches(4.5))
    _kill_line(tb)
    tf = tb.text_frame; tf.word_wrap = True
    for i, p_text in enumerate(punti):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f" ·  {p_text}"
        p.font.size = Pt(20)
        p.font.color.rgb = BIANCO
        p.font.name = F_UI
        p.space_after = Pt(10)
        _no_proof(p)
    return s


def agenda_slide(prs, items):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, PAPER)

    _strip_top(s, f"Agenda  ·  {len(items):02d} punti", dot_color=COBALT)

    # Titolo grande in italiano
    _tb(s, ML, 1.1, 5, 0.9, "Il",
        sz=72, bold=True, col=INK, font=F_DISP)
    _tb(s, ML, 1.8, 5, 0.9, "programma.",
        sz=72, col=TANGERINE, font=F_DISP, italic=True)

    # Sottotesto
    _tb(s, ML, 2.75, 4.5, 0.6,
        "Cosa vedremo in questo modulo.",
        sz=18, col=INK_DIM, font=F_UI)

    # Card numerate a destra
    x_cards = 6.0
    card_w = 6.2
    y_start = 1.1
    card_h = max(0.52, min(0.68, 5.6 / len(items)))

    for i, item in enumerate(items):
        y = y_start + i * (card_h + 0.07)
        if y + card_h > 6.7:
            break
        _rect(s, x_cards, y, card_w, card_h, BIANCO,
              line_color=RULE, radius=True)
        _tb(s, x_cards + 0.2, y + 0.1, 0.45, card_h - 0.2,
            f"{i+1:02d}", sz=14, col=TANGERINE, font=F_MONO, bold=True)
        _tb(s, x_cards + 0.7, y + 0.1, card_w - 1.0, card_h - 0.2,
            item, sz=16, col=INK, font=F_DISP)

    return s


def question_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, INK)

    _tb(s, ML, 2.2, CW, 1.5, "Domande?",
        sz=96, bold=True, col=BIANCO, font=F_DISP, italic=True,
        align=PP_ALIGN.CENTER)

    _rect(s, 5.5, 4.0, 2.3, 0.05, LIME)

    _tb(s, ML, 4.4, CW, 0.8,
        "Non esistono domande stupide.",
        sz=22, col=INK_FAINT, font=F_UI, align=PP_ALIGN.CENTER)

    _tb(s, ML, Y_STRIP_BOT, CW, 0.3, "Q & A",
        sz=12, col=INK_DIM, font=F_MONO)
    return s


def next_module_slide(prs, testo):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, COBALT)

    _tb(s, ML, 1.5, CW, 0.4, "PROSSIMO MODULO",
        sz=13, col=LIME, font=F_MONO, bold=True)

    _tb(s, ML, 2.5, CW, 2.5, testo,
        sz=40, bold=True, col=BIANCO, font=F_DISP)

    _rect(s, ML, 5.2, 1.5, 0.05, LIME)
    return s


def salva(prs, path):
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        is_dark = False
        try:
            rgb = slide.background.fill.fore_color.rgb
            if rgb in (INK, COBALT):
                is_dark = True
        except Exception:
            pass

        num_col = BIANCO if is_dark else INK
        tot_col = INK_DIM if not is_dark else INK_FAINT

        _tb(slide, 10.5, Y_STRIP_BOT, 1.0, 0.35,
            f"{i:02d}", sz=20, col=num_col, font=F_DISP, bold=True,
            align=PP_ALIGN.RIGHT)
        _tb(slide, 11.5, Y_STRIP_BOT, 1.6, 0.35,
            f"/ {total}", sz=20, col=tot_col, font=F_DISP,
            align=PP_ALIGN.LEFT)

    prs.save(path)
    print(f"Salvato: {path}")
