#!/usr/bin/env python3
"""
Genera il template del certificato di completamento corso AI Basics.
Stile coerente con le slide: paper bg, Georgia serif, accenti vibranti.
Due versioni: IT e EN.
"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# ── Palette (stessa delle slide) ────────────────────────
PAPER     = RGBColor(0xFA, 0xFA, 0xF7)
INK       = RGBColor(0x0B, 0x0B, 0x0C)
INK_2     = RGBColor(0x2A, 0x2A, 0x2D)
INK_DIM   = RGBColor(0x6B, 0x6B, 0x72)
INK_FAINT = RGBColor(0xB5, 0xB5, 0xBC)
BIANCO    = RGBColor(0xFF, 0xFF, 0xFF)
LIME      = RGBColor(0xC6, 0xFF, 0x3D)
TANGERINE = RGBColor(0xFF, 0x5A, 0x1F)
COBALT    = RGBColor(0x1E, 0x3C, 0xFF)
MAGENTA   = RGBColor(0xFF, 0x2E, 0x88)
AMBER     = RGBColor(0xFF, 0xB8, 0x00)
MINT      = RGBColor(0x00, 0xC2, 0x8A)
RULE      = RGBColor(0xE6, 0xE5, 0xDF)

F_DISP = "Georgia"
F_UI   = "Calibri"
F_MONO = "Consolas"

SW = Inches(13.333)
SH = Inches(7.5)


def _kill_line(shape):
    spPr = shape._element.spPr
    ln = spPr.find(qn('a:ln'))
    if ln is None:
        ln = etree.SubElement(spPr, qn('a:ln'))
    else:
        for child in list(ln):
            ln.remove(child)
    ln.set('w', '0')
    etree.SubElement(ln, qn('a:noFill'))


def _rect(slide, l, t, w, h, color, line_color=None, line_w=None, radius=None):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(st, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color
    if line_color:
        s.line.color.rgb = line_color; s.line.width = Pt(line_w or 1)
    else:
        _kill_line(s)
    return s


def _tb(slide, l, t, w, h, txt, sz=18, bold=False, col=INK,
        align=PP_ALIGN.LEFT, font=F_UI, italic=False):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    _kill_line(tb)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = txt; p.font.size = Pt(sz); p.font.bold = bold
    p.font.color.rgb = col; p.font.name = font
    p.font.italic = italic; p.alignment = align
    for run in p.runs:
        rPr = run._r.find(qn('a:rPr'))
        if rPr is not None:
            rPr.set('noProof', '1')
    return tf


def _dot(slide, l, t, color, size=0.1):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                               Inches(l), Inches(t), Inches(size), Inches(size))
    c.fill.solid(); c.fill.fore_color.rgb = color; _kill_line(c)


def genera_certificato(lang="it"):
    prs = Presentation()
    prs.slide_width = SW; prs.slide_height = SH
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # ── Sfondo paper ────────────────────────────────────
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = PAPER

    # ── Cornice decorativa esterna (bordo sottile) ──────
    _rect(s, 0.35, 0.25, 12.62, 7.0, PAPER, line_color=RULE, line_w=1.5)

    # ── Barra colorata superiore (5 accenti affiancati) ─
    bar_colors = [LIME, TANGERINE, COBALT, MAGENTA, MINT]
    bar_w = 12.62 / len(bar_colors)
    for i, c in enumerate(bar_colors):
        _rect(s, 0.35 + i * bar_w, 0.25, bar_w + 0.01, 0.12, c)

    # ── Barra colorata inferiore ────────────────────────
    for i, c in enumerate(bar_colors):
        _rect(s, 0.35 + i * bar_w, 7.13, bar_w + 0.01, 0.12, c)

    # ── Dot decorativi angoli ───────────────────────────
    _dot(s, 0.6, 0.6, TANGERINE, 0.12)
    _dot(s, 12.6, 0.6, COBALT, 0.12)
    _dot(s, 0.6, 6.8, MAGENTA, 0.12)
    _dot(s, 12.6, 6.8, MINT, 0.12)

    # ── Contenuti (variano per lingua) ──────────────────
    if lang == "it":
        eyebrow = "AI BASICS  ·  CERTIFICATO DI COMPLETAMENTO"
        pre_title = "Si attesta che"
        participant = "[NOME E COGNOME]"
        post_title = "ha completato con successo il corso"
        course = "AI Basics: Strumenti di Intelligenza\nArtificiale per Tutti"
        duration = "Durata: 20 ore  ·  8 moduli  ·  Teoria + Pratica"
        skills_label = "COMPETENZE CERTIFICATE"
        skills = [
            "Uso professionale di ChatGPT, Claude e modelli open source",
            "Prompt engineering con tecniche strutturate",
            "Integrazione dell'AI nei workflow lavorativi",
            "Uso responsabile e consapevole dell'AI",
        ]
        score_label = "VALUTAZIONE"
        score_fields = "Progetto: ___ / 25     Quiz: ___ / 20"
        date_label = "Data"
        date_val = "____________________"
        sign_label = "Firma del Formatore"
        sign_val = "____________________"
    else:
        eyebrow = "AI BASICS  ·  CERTIFICATE OF COMPLETION"
        pre_title = "This certifies that"
        participant = "[FULL NAME]"
        post_title = "has successfully completed the course"
        course = "AI Basics: Artificial Intelligence\nTools for Everyone"
        duration = "Duration: 20 hours  ·  8 modules  ·  Theory + Practice"
        skills_label = "CERTIFIED SKILLS"
        skills = [
            "Professional use of ChatGPT, Claude and open source models",
            "Prompt engineering with structured techniques",
            "AI integration into daily work workflows",
            "Responsible and informed use of AI",
        ]
        score_label = "ASSESSMENT"
        score_fields = "Project: ___ / 25     Quiz: ___ / 20"
        date_label = "Date"
        date_val = "____________________"
        sign_label = "Instructor Signature"
        sign_val = "____________________"

    # ── Eyebrow ─────────────────────────────────────────
    _dot(s, 1.0, 0.75, TANGERINE, 0.08)
    _tb(s, 1.2, 0.7, 11, 0.3, eyebrow,
        sz=11, col=INK_DIM, font=F_MONO, bold=True, align=PP_ALIGN.CENTER)

    # ── Pre-titolo ──────────────────────────────────────
    _tb(s, 1.0, 1.2, 11.3, 0.4, pre_title,
        sz=18, col=INK_DIM, font=F_UI, align=PP_ALIGN.CENTER)

    # ── Nome partecipante (grande, serif, accento) ──────
    _tb(s, 1.0, 1.6, 11.3, 0.9, participant,
        sz=52, bold=True, col=COBALT, font=F_DISP, italic=True,
        align=PP_ALIGN.CENTER)

    # ── Linea decorativa sotto il nome ──────────────────
    _rect(s, 4.5, 2.55, 4.3, 0.04, TANGERINE)

    # ── Post-titolo ─────────────────────────────────────
    _tb(s, 1.0, 2.75, 11.3, 0.4, post_title,
        sz=18, col=INK_DIM, font=F_UI, align=PP_ALIGN.CENTER)

    # ── Nome corso (grande, serif) ──────────────────────
    _tb(s, 1.0, 3.15, 11.3, 1.2, course,
        sz=36, bold=True, col=INK, font=F_DISP, align=PP_ALIGN.CENTER)

    # ── Durata ──────────────────────────────────────────
    _tb(s, 1.0, 4.2, 11.3, 0.35, duration,
        sz=14, col=INK_DIM, font=F_MONO, align=PP_ALIGN.CENTER)

    # ── Card competenze ─────────────────────────────────
    card_x, card_y, card_w, card_h = 1.8, 4.7, 9.7, 1.5
    _rect(s, card_x, card_y, card_w, card_h, BIANCO,
          line_color=RULE, line_w=0.75, radius=True)

    _tb(s, card_x + 0.3, card_y + 0.15, card_w - 0.6, 0.25,
        skills_label, sz=10, col=INK_DIM, font=F_MONO, bold=True)

    # Linea sotto label competenze
    _rect(s, card_x + 0.3, card_y + 0.4, 1.2, 0.025, LIME)

    # 4 pill di competenza in griglia 2x2
    pill_colors = [LIME, TANGERINE, COBALT, MAGENTA]
    for i, (skill, pc) in enumerate(zip(skills, pill_colors)):
        col_idx = i % 2
        row_idx = i // 2
        px = card_x + 0.3 + col_idx * 4.8
        py = card_y + 0.55 + row_idx * 0.45

        # Dot colorato
        _dot(s, px, py + 0.05, pc, 0.07)
        # Testo competenza
        _tb(s, px + 0.15, py, 4.4, 0.35, skill,
            sz=12, col=INK_2, font=F_UI)

    # ── Valutazione ─────────────────────────────────────
    _tb(s, 1.0, 6.35, 4, 0.2, score_label,
        sz=10, col=INK_DIM, font=F_MONO, bold=True)
    _tb(s, 1.0, 6.55, 5, 0.25, score_fields,
        sz=14, col=INK_2, font=F_MONO)

    # ── Data e firma ────────────────────────────────────
    # Data (centro)
    _tb(s, 5.5, 6.35, 3, 0.2, date_label,
        sz=10, col=INK_DIM, font=F_MONO, bold=True, align=PP_ALIGN.CENTER)
    _tb(s, 5.5, 6.55, 3, 0.25, date_val,
        sz=14, col=INK_2, font=F_UI, align=PP_ALIGN.CENTER)

    # Firma (destra)
    _tb(s, 9.0, 6.35, 3.9, 0.2, sign_label,
        sz=10, col=INK_DIM, font=F_MONO, bold=True, align=PP_ALIGN.CENTER)
    _tb(s, 9.0, 6.55, 3.9, 0.25, sign_val,
        sz=14, col=INK_2, font=F_UI, align=PP_ALIGN.CENTER)

    return prs


# ── Genera entrambe le versioni ─────────────────────────
if __name__ == "__main__":
    base = os.path.dirname(os.path.abspath(__file__))

    # IT
    prs_it = genera_certificato("it")
    out_it = os.path.join(base, "materiale_trasversale", "Certificato_AI_Basics.pptx")
    prs_it.save(out_it)
    print(f"  OK  {out_it}")

    # EN
    prs_en = genera_certificato("en")
    out_en = os.path.join(base, "..", "ai_course_basics_en", "cross_cutting",
                          "Certificate_AI_Basics.pptx")
    os.makedirs(os.path.dirname(out_en), exist_ok=True)
    prs_en.save(out_en)
    print(f"  OK  {out_en}")

    print("  Certificati generati!")
