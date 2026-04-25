#!/usr/bin/env python3
"""
Certificato v3 — "Glacier": sfondo bianco ghiaccio, accenti blu/teal sfumati,
linee geometriche sottili, stile nordico-minimal.
Genera versione IT e EN.
"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# ── Palette "Glacier" ───────────────────────────────────
ICE       = RGBColor(0xF4, 0xF7, 0xFB)   # sfondo ghiaccio
ICE_2     = RGBColor(0xE8, 0xEE, 0xF5)   # card sfondo
SLATE     = RGBColor(0x1A, 0x1F, 0x36)   # testo principale
SLATE_2   = RGBColor(0x3B, 0x42, 0x5A)   # testo secondario
SLATE_DIM = RGBColor(0x8B, 0x92, 0xA8)   # testo dim
OCEAN     = RGBColor(0x0A, 0x6E, 0xBD)   # blu oceano
SKY       = RGBColor(0x38, 0xB2, 0xEB)   # azzurro cielo
TEAL_G    = RGBColor(0x0D, 0xB4, 0x9B)   # teal glaciale
FROST     = RGBColor(0xCB, 0xDF, 0xF0)   # bordi frost
BIANCO    = RGBColor(0xFF, 0xFF, 0xFF)

F_DISP = "Georgia"
F_UI   = "Calibri"
F_MONO = "Consolas"

SW = Inches(13.333)
SH = Inches(7.5)


def _kill_line(shape):
    spPr = shape._element.spPr
    ln = spPr.find(qn('a:ln'))
    if ln is None:
        ln = etree.SubElement(spPr, qn('a:ln'))
    else:
        for child in list(ln):
            ln.remove(child)
    ln.set('w', '0')
    etree.SubElement(ln, qn('a:noFill'))


def _rect(slide, l, t, w, h, color, line_color=None, line_w=None, radius=None):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(st, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color
    if line_color:
        s.line.color.rgb = line_color; s.line.width = Pt(line_w or 1)
    else:
        _kill_line(s)
    return s


def _tb(slide, l, t, w, h, txt, sz=18, bold=False, col=SLATE,
        align=PP_ALIGN.LEFT, font=F_UI, italic=False):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    _kill_line(tb)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = txt; p.font.size = Pt(sz); p.font.bold = bold
    p.font.color.rgb = col; p.font.name = font
    p.font.italic = italic; p.alignment = align
    for run in p.runs:
        rPr = run._r.find(qn('a:rPr'))
        if rPr is not None:
            rPr.set('noProof', '1')
    return tf


def _circle(slide, cx, cy, size, color):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL,
        Inches(cx - size/2), Inches(cy - size/2), Inches(size), Inches(size))
    s.fill.solid(); s.fill.fore_color.rgb = color; _kill_line(s)


def _hex(slide, cx, cy, size, color):
    """Esagono decorativo."""
    s = slide.shapes.add_shape(MSO_SHAPE.HEXAGON,
        Inches(cx - size/2), Inches(cy - size/2), Inches(size), Inches(size))
    s.fill.solid(); s.fill.fore_color.rgb = color; _kill_line(s)


def genera_certificato_v3(lang="it"):
    prs = Presentation()
    prs.slide_width = SW; prs.slide_height = SH
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # ── Sfondo ghiaccio ─────────────────────────────────
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = ICE

    # ── Banda laterale sinistra (gradient simulato 3 strisce) ──
    _rect(s, 0, 0, 0.12, 7.5, OCEAN)
    _rect(s, 0.12, 0, 0.06, 7.5, SKY)
    _rect(s, 0.18, 0, 0.04, 7.5, TEAL_G)

    # ── Linea sottile superiore ─────────────────────────
    _rect(s, 0.5, 0.55, 12.5, 0.012, FROST)

    # ── Linea sottile inferiore ─────────────────────────
    _rect(s, 0.5, 6.95, 12.5, 0.012, FROST)

    # ── Cerchi decorativi sfumati (angolo alto-destra) ──
    _circle(s, 12.2, 0.9, 0.5, ICE_2)
    _circle(s, 12.45, 0.65, 0.25, FROST)
    _circle(s, 11.8, 0.55, 0.15, FROST)

    # ── Esagoni decorativi (angolo basso-destra) ────────
    _hex(s, 12.3, 6.5, 0.35, ICE_2)
    _hex(s, 11.9, 6.7, 0.2, FROST)

    # ── Contenuti ───────────────────────────────────────
    if lang == "it":
        eyebrow = "AI BASICS  ·  CERTIFICATO DI COMPLETAMENTO"
        pre_title = "Si attesta che"
        participant = "[NOME E COGNOME]"
        post_title = "ha completato con successo il corso"
        course = "AI Basics"
        course_sub = "Strumenti di Intelligenza Artificiale per Tutti"
        duration = "20 ore  ·  8 moduli  ·  Teoria + Pratica"
        skills_label = "COMPETENZE CERTIFICATE"
        skills = [
            "Uso professionale di ChatGPT, Claude e modelli open source",
            "Prompt engineering con tecniche strutturate",
            "Integrazione dell'AI nei workflow lavorativi",
            "Uso responsabile e consapevole dell'AI",
        ]
        score_label = "VALUTAZIONE"
        score_fields = "Progetto: ___ / 25     Quiz: ___ / 20"
        date_label = "Data"
        date_val = "____________________"
        sign_label = "Firma del Formatore"
        sign_val = "____________________"
    else:
        eyebrow = "AI BASICS  ·  CERTIFICATE OF COMPLETION"
        pre_title = "This certifies that"
        participant = "[FULL NAME]"
        post_title = "has successfully completed the course"
        course = "AI Basics"
        course_sub = "Artificial Intelligence Tools for Everyone"
        duration = "20 hours  ·  8 modules  ·  Theory + Practice"
        skills_label = "CERTIFIED SKILLS"
        skills = [
            "Professional use of ChatGPT, Claude and open source models",
            "Prompt engineering with structured techniques",
            "AI integration into daily work workflows",
            "Responsible and informed use of AI",
        ]
        score_label = "ASSESSMENT"
        score_fields = "Project: ___ / 25     Quiz: ___ / 20"
        date_label = "Date"
        date_val = "____________________"
        sign_label = "Instructor Signature"
        sign_val = "____________________"

    # ── Eyebrow ─────────────────────────────────────────
    _tb(s, 0.8, 0.7, 11.5, 0.3, eyebrow,
        sz=11, col=OCEAN, font=F_MONO, bold=True, align=PP_ALIGN.CENTER)

    # ── Pre-titolo ──────────────────────────────────────
    _tb(s, 0.8, 1.3, 11.5, 0.35, pre_title,
        sz=20, col=SLATE_DIM, font=F_UI, align=PP_ALIGN.CENTER)

    # ── Nome partecipante ───────────────────────────────
    _tb(s, 0.8, 1.75, 11.5, 0.9, participant,
        sz=56, bold=True, col=OCEAN, font=F_DISP, italic=True,
        align=PP_ALIGN.CENTER)

    # ── Linea teal sotto il nome ────────────────────────
    _rect(s, 4.0, 2.7, 5.3, 0.035, TEAL_G)
    # Due cerchietti ai lati della linea
    _circle(s, 3.9, 2.72, 0.07, SKY)
    _circle(s, 9.4, 2.72, 0.07, SKY)

    # ── Post-titolo ─────────────────────────────────────
    _tb(s, 0.8, 2.9, 11.5, 0.35, post_title,
        sz=20, col=SLATE_DIM, font=F_UI, align=PP_ALIGN.CENTER)

    # ── Nome corso ──────────────────────────────────────
    _tb(s, 0.8, 3.3, 11.5, 0.7, course,
        sz=44, bold=True, col=SLATE, font=F_DISP, align=PP_ALIGN.CENTER)

    # ── Sottotitolo corso ───────────────────────────────
    _tb(s, 0.8, 3.9, 11.5, 0.5, course_sub,
        sz=22, col=SLATE_2, font=F_DISP, italic=True, align=PP_ALIGN.CENTER)

    # ── Durata (pill arrotondato) ───────────────────────
    pill_w = 5.5
    pill_x = (13.333 - pill_w) / 2
    _rect(s, pill_x, 4.4, pill_w, 0.35, ICE_2, line_color=FROST, line_w=0.75,
          radius=True)
    _tb(s, pill_x, 4.42, pill_w, 0.3, duration,
        sz=12, col=SLATE_DIM, font=F_MONO, align=PP_ALIGN.CENTER)

    # ── Card competenze ─────────────────────────────────
    card_x, card_y, card_w, card_h = 1.5, 5.0, 10.3, 1.55
    _rect(s, card_x, card_y, card_w, card_h, BIANCO,
          line_color=FROST, line_w=0.75, radius=True)

    _tb(s, card_x + 0.4, card_y + 0.12, card_w - 0.8, 0.22,
        skills_label, sz=10, col=OCEAN, font=F_MONO, bold=True)

    # Linea sotto label
    _rect(s, card_x + 0.4, card_y + 0.38, 1.0, 0.025, TEAL_G)

    # 4 skill in griglia 2x2 con icone cerchio
    skill_colors = [OCEAN, SKY, TEAL_G, OCEAN]
    for i, (skill, sc) in enumerate(zip(skills, skill_colors)):
        col_idx = i % 2
        row_idx = i // 2
        px = card_x + 0.4 + col_idx * 5.0
        py = card_y + 0.5 + row_idx * 0.45

        # Cerchio colorato
        _circle(s, px + 0.05, py + 0.12, 0.09, sc)
        # Testo
        _tb(s, px + 0.18, py, 4.5, 0.35, skill,
            sz=12, col=SLATE_2, font=F_UI)

    # ── Sezione inferiore: Valutazione + Data + Firma ───
    _tb(s, 1.0, 6.6, 3.5, 0.15, score_label,
        sz=9, col=OCEAN, font=F_MONO, bold=True)
    _tb(s, 1.0, 6.75, 3.5, 0.2, score_fields,
        sz=12, col=SLATE_2, font=F_MONO)

    _tb(s, 5.2, 6.6, 3, 0.15, date_label,
        sz=9, col=OCEAN, font=F_MONO, bold=True, align=PP_ALIGN.CENTER)
    _tb(s, 5.2, 6.75, 3, 0.2, date_val,
        sz=13, col=SLATE_2, font=F_UI, align=PP_ALIGN.CENTER)

    _tb(s, 8.8, 6.6, 3.9, 0.15, sign_label,
        sz=9, col=OCEAN, font=F_MONO, bold=True, align=PP_ALIGN.CENTER)
    _tb(s, 8.8, 6.75, 3.9, 0.2, sign_val,
        sz=13, col=SLATE_2, font=F_UI, align=PP_ALIGN.CENTER)

    return prs


# ── Genera entrambe le versioni ─────────────────────────
if __name__ == "__main__":
    base = os.path.dirname(os.path.abspath(__file__))

    prs_it = genera_certificato_v3("it")
    out_it = os.path.join(base, "materiale_trasversale",
                          "Certificato_AI_Basics_Glacier.pptx")
    prs_it.save(out_it)
    print(f"  OK  {out_it}")

    prs_en = genera_certificato_v3("en")
    out_en = os.path.join(base, "..", "ai_course_basics_en", "cross_cutting",
                          "Certificate_AI_Basics_Glacier.pptx")
    os.makedirs(os.path.dirname(out_en), exist_ok=True)
    prs_en.save(out_en)
    print(f"  OK  {out_en}")

    print("  Certificati Glacier generati!")
